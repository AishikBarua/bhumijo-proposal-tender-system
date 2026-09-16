#!/usr/bin/env python3
"""
Export everything in the database to files you can actually open.

bhumijo.db is a binary database file — Windows has no application for it, so
double-clicking does nothing. This writes the same data as spreadsheets.

    python tools/export_data.py

Produces, in data/export/:
    Bhumijo_Data_<date>.xlsx   one sheet per table  (if openpyxl is installed)
    *.csv                       one file per table  (always — Excel opens these)

It only reads. Nothing in the database is changed, and it is safe to run
while the server is running and people are using the tracker.
"""

from __future__ import annotations

import csv
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backend.config import settings  # noqa: E402

# Friendly sheet names, and the order a person would want them in.
TABLES = [
    ("proposals", "Proposals"),
    ("grant_programmes", "Grant programmes"),
    ("clients", "Clients"),
    ("audit_log", "Change history"),
    ("categories", "Categories"),
    ("entities", "Entities"),
    ("users", "Users"),
    ("roles", "Roles"),
]


def read_table(conn: sqlite3.Connection, table: str) -> tuple[list[str], list[tuple]]:
    cursor = conn.execute(f"SELECT * FROM {table}")
    columns = [d[0] for d in cursor.description]
    return columns, cursor.fetchall()


def write_csvs(conn: sqlite3.Connection, out_dir: Path) -> list[Path]:
    written = []
    for table, label in TABLES:
        try:
            columns, rows = read_table(conn, table)
        except sqlite3.OperationalError:
            continue
        path = out_dir / f"{label.replace(' ', '_')}.csv"
        # utf-8-sig so Excel shows accented and Bengali text correctly rather
        # than mojibake — plain utf-8 is misread by Excel on Windows.
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(columns)
            writer.writerows(rows)
        written.append(path)
        print(f"  {label:20} {len(rows):>5} rows  ->  {path.name}")
    return written


def write_xlsx(conn: sqlite3.Connection, out_dir: Path) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("\n  (openpyxl is not installed, so no .xlsx was made — the CSV")
        print("   files above open in Excel just as well. To get a single")
        print("   workbook instead, run:  python -m pip install openpyxl)")
        return None

    workbook = Workbook()
    workbook.remove(workbook.active)

    header_fill = PatternFill("solid", fgColor="1B4F86")
    header_font = Font(color="FFFFFF", bold=True)

    for table, label in TABLES:
        try:
            columns, rows = read_table(conn, table)
        except sqlite3.OperationalError:
            continue

        sheet = workbook.create_sheet(label[:31])
        sheet.append(columns)
        for row in rows:
            sheet.append(["" if v is None else v for v in row])

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")

        # Readable column widths, capped so a long remark does not make one
        # column fill the screen.
        for index, column in enumerate(columns, start=1):
            longest = max([len(str(column))] +
                          [len(str(r[index - 1])) for r in rows[:200]] or [10])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 55)

        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions

    path = out_dir / f"Bhumijo_Data_{datetime.now():%Y%m%d}.xlsx"
    workbook.save(path)
    return path


def main() -> int:
    database = settings.database_path
    if not database.exists():
        print(f"No database found at {database}")
        print("Run SETUP_AND_START.bat first.")
        return 1

    out_dir = settings.data_dir / "export"
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Reading {database}")
    print(f"Writing to {out_dir}\n")

    # Read-only, so this cannot disturb a running server.
    conn = sqlite3.connect(f"file:{database}?mode=ro", uri=True)
    try:
        write_csvs(conn, out_dir)
        workbook = write_xlsx(conn, out_dir)
    finally:
        conn.close()

    print()
    if workbook:
        print(f"  Excel workbook: {workbook.name}")
    print(f"\nDone. Everything is in:\n  {out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
